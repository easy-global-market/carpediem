#!/usr/bin/env python3
"""Build the Carpediem V2 activity-pressure matrix (matrice_act_press.ods) for the
Mer du Nord study, from the official (January 2026) matrix.

Rules: see migration_v1_v2/README.md; decision numbers in comments (D1...) refer to the
study's decision register (not versioned).

Inputs
  --param     official parameter file (sheet matrice_activites_pressions)
  --workbook  correspondence workbook (sheets Pressions, Activités), see
              ref/modele_correspondances_v1_v2.xlsx
  --template  V2 reference matrix (Baie de Seine), source of the V2 typologies
Outputs
  in --out-dir (the config/ folder of the V2 study):
    matrice_act_press.ods      V2 matrix: CCR_mediane, CCR_precaution, CCR_binaire,
                               typologie_activites, typologie_pressions
  in --report-dir (default: --out-dir):
    conversion_activites.csv   origin of every V2 activity row
    conversion_pressions.csv   official -> V2 pressure correspondence
    conversion_log.txt         counts and checks
"""
import argparse
import os
import sys

import numpy as np
import openpyxl
import pandas as pd

IC_SCALE = 5.0          # D4: IC V2 = IC official / 5
BINARY_IC_MIN = 3       # D6: binary link = 1 if amplitude > 0 and IC official >= 3
AMPLITUDE_SCALE = 1000  # D1: CCR = amplitude / 1000


def read_official(path):
    raw = pd.read_excel(path, engine="odf", sheet_name="matrice_activites_pressions", header=None)
    codes_p = raw.iloc[0, 4:].tolist()
    body = raw.iloc[2:].copy()
    body.columns = ["Type", "Code_A", "Activity", "Phase"] + codes_p
    amp = body[body.Type == "Final"].set_index("Code_A")[codes_p].apply(pd.to_numeric, errors="coerce")
    ic = body[body.Type == "IC"].set_index("Code_A")[codes_p].apply(pd.to_numeric, errors="coerce")
    names = body[body.Type == "Final"].set_index("Code_A")[["Activity", "Phase"]]
    return amp, ic, names


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    # Pressions: A code officiel, D code V2 proposé, H évaluer, I décision, J code retenu
    press = []
    for r in wb["Pressions"].iter_rows(min_row=3, values_only=True):
        if not r[0]:
            continue
        press.append(dict(official=r[0], label=r[1], v2=(r[9] or r[3]), evaluate=(r[7] or "").strip(),
                          decision=(r[8] or "").strip()))
    # Activités: A couche, C code officiel, F code V2 proposé, J décision, K code retenu, M traitement
    acts = []
    for r in wb["Activités"].iter_rows(min_row=3, values_only=True):
        if not r[0]:
            continue
        acts.append(dict(layer=r[0].replace(" ", ""), column=r[1], official=r[2], v2=(r[10] or r[5]),
                         decision=(r[9] or "").strip(), multi=(r[12] or "")))
    return press, acts


def parent_code(code):
    parts = code.split("_")
    return "_".join(parts[:-1]) if len(parts) > 2 else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--param", required=True)
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--report-dir")
    a = ap.parse_args()
    a.report_dir = a.report_dir or a.out_dir
    os.makedirs(a.out_dir, exist_ok=True)
    os.makedirs(a.report_dir, exist_ok=True)
    log = []

    amp, ic, off_names = read_official(a.param)
    press, acts = read_workbook(a.workbook)
    typo_act = pd.read_excel(a.template, engine="odf", sheet_name="typologie_activites")
    typo_press = pd.read_excel(a.template, engine="odf", sheet_name="typologie_pressions")

    # ---- Pressures (D11 to D15) ---------------------------------------------------
    evaluated = [p for p in press if p["evaluate"] == "Oui" and p["v2"]]
    p_map = {}  # V2 code -> official code (one official pressure per evaluated V2 leaf)
    for p in evaluated:
        if p["v2"] in p_map:
            sys.exit(f"Two evaluated official pressures for {p['v2']}")
        p_map[p["v2"]] = p["official"]
    known_p = set(typo_press.Code_P)
    unknown = [c for c in p_map if c not in known_p]
    if unknown:
        sys.exit(f"Unknown V2 pressure codes: {unknown}")
    missing_cols = [o for o in p_map.values() if o not in amp.columns]
    if missing_cols:
        sys.exit(f"Official pressures missing from the matrix: {missing_cols}")
    # keep the order of the V2 typology
    pcols = [c for c in typo_press.Code_P if c in p_map]
    log.append(f"Evaluated pressures: {len(pcols)}")

    pd.DataFrame([dict(code_officiel=p["official"], libelle_officiel=p["label"], code_v2=p["v2"] or "",
                       evaluee=p["evaluate"], decision=p["decision"]) for p in press]) \
        .to_csv(os.path.join(a.report_dir, "conversion_pressions.csv"), index=False)

    # demonstrateur = 1 for evaluated leaves and their ancestors (D12)
    parent_of = dict(zip(typo_press.Code_P, typo_press.parent))
    demo = set()
    for c in pcols:
        while isinstance(c, str) and c:
            demo.add(c)
            c = parent_of.get(c)
    typo_press_out = typo_press.copy()
    typo_press_out["demonstrateur"] = typo_press_out.Code_P.isin(demo).astype(int)

    # ---- Activities (D16 to D21) -------------------------------------------------
    # Per data layer: retained V2 code = code of the first row of the layer (D17, D19);
    # official activities kept = the rows of the layer with that code (merged with D18).
    kept = [x for x in acts if x["decision"] in ("Validé", "Modifié")]
    layer_code = {}
    for x in kept:
        layer_code.setdefault(x["layer"], x["v2"])
    direct = {}   # V2 code -> list of official codes
    origin = {}   # V2 code -> (layers, rule)
    skipped = []
    for x in kept:
        if x["v2"] != layer_code[x["layer"]]:
            skipped.append(f"{x['layer']} {x['official']} -> {x['v2']} not retained (one code per layer, D17/D19)")
            continue
        if not x["official"]:
            skipped.append(f"{x['layer']} -> {x['v2']}: no official activity, no amplitude (D20, O1)")
            continue
        direct.setdefault(x["v2"], [])
        if x["official"] not in direct[x["v2"]]:
            direct[x["v2"]].append(x["official"])
        origin.setdefault(x["v2"], set()).add(x["layer"])
    log += skipped
    known_a = set(typo_act.Code_A)
    unknown = [c for c in direct if c not in known_a]
    if unknown:
        sys.exit(f"Unknown V2 activity codes: {unknown}")

    # Leaf values: CCR (median), IC and binary link per V2 code with a direct mapping
    def convert(codes):
        """D1 to D6, D18 (max over the merged official activities)."""
        A = amp.loc[codes, [p_map[c] for c in pcols]].to_numpy(dtype=float)
        I = ic.loc[codes, [p_map[c] for c in pcols]].to_numpy(dtype=float)
        A = np.nan_to_num(A, nan=0.0)          # D3
        I = np.nan_to_num(I, nan=0.0)          # D5
        lien = np.zeros(len(pcols)); conf = np.zeros(len(pcols)); binl = np.zeros(len(pcols))
        for j in range(len(pcols)):
            # D18: activity with the max amplitude; on a tie, the lowest IC
            order = sorted(range(len(codes)), key=lambda i: (-A[i, j], I[i, j]))
            i = order[0]
            lien[j] = A[i, j] / AMPLITUDE_SCALE                              # D1
            conf[j] = I[i, j] / IC_SCALE                                     # D4
            binl[j] = 1.0 if (A[i, j] > 0 and I[i, j] >= BINARY_IC_MIN) else 0.0  # D6
        return lien, conf, binl

    vals = {}
    for code, offs in direct.items():
        vals[code] = convert(offs)

    # Parent levels (D8): max of the children, IC of the child carrying the max
    # (lowest on a tie), binary = max. A code with a direct mapping (e.g. A_7_1)
    # takes part in the max with its own values.
    codes_sorted = sorted(typo_act.Code_A, key=lambda c: -c.count("_"))
    children = {}
    for c in typo_act.Code_A:
        p = parent_code(c)
        if p is not None:
            children.setdefault(p, []).append(c)
    zero = (np.zeros(len(pcols)),) * 3
    for c in codes_sorted:
        cands = [vals[k] for k in children.get(c, []) if k in vals]
        if c in vals:
            cands.append(vals[c])
        if not cands:
            continue
        L = np.array([x[0] for x in cands]); I = np.array([x[1] for x in cands]); B = np.array([x[2] for x in cands])
        lien = L.max(axis=0); conf = np.zeros(len(pcols))
        for j in range(len(pcols)):
            idx = [i for i in range(len(cands)) if L[i, j] == lien[j]]
            conf[j] = min(I[i, j] for i in idx)
        vals[c] = (lien, conf, B.max(axis=0))

    # ---- Sheets ---------------------------------------------------------------
    def sheet(kind):
        rows = []
        for _, r in typo_act.iterrows():
            lien, conf, binl = vals.get(r.Code_A, zero)   # D9: no correspondence -> 0
            l = binl if kind == "binaire" else lien
            rows.append([r.Niveau, r.Code_A, r.Nom_activite, "lien"] + list(l))
            rows.append([r.Niveau, r.Code_A, r.Nom_activite, "ic"] + list(conf))
        return pd.DataFrame(rows, columns=["Niveau", "Code_A", "Nom_activite", "Type"] + pcols)

    med = sheet("mediane")
    out = os.path.join(a.out_dir, "matrice_act_press.ods")
    with pd.ExcelWriter(out, engine="odf") as w:
        med.to_excel(w, sheet_name="CCR_mediane", index=False)
        med.to_excel(w, sheet_name="CCR_precaution", index=False)   # D7
        sheet("binaire").to_excel(w, sheet_name="CCR_binaire", index=False)
        typo_act.to_excel(w, sheet_name="typologie_activites", index=False)
        typo_press_out.to_excel(w, sheet_name="typologie_pressions", index=False)

    # ---- Conversion report ------------------------------------------------------
    rep = []
    for _, r in typo_act.iterrows():
        c = r.Code_A
        if c in direct:
            rule = "direct" + (" (max of merged activities, D18)" if len(direct[c]) > 1 else "")
            if children.get(c) and any(k in vals for k in children[c]):
                rule += " + children (D8)"
            src = " + ".join(direct[c])
        elif c in vals:
            rule, src = "aggregated from children (D8)", ""
        else:
            rule, src = "no correspondence, 0 (D9)", ""
        rep.append(dict(Niveau=r.Niveau, Code_A=c, Nom_activite=r.Nom_activite, regle=rule,
                        activites_officielles=src,
                        couches=" ; ".join(sorted(origin.get(c, []))),
                        nb_pressions_liees=int((vals.get(c, zero)[0] > 0).sum())))
    pd.DataFrame(rep).to_csv(os.path.join(a.report_dir, "conversion_activites.csv"), index=False)

    log.append(f"V2 activities: {len(typo_act)}; with direct mapping: {len(direct)}; "
               f"with values (direct or aggregated): {len(vals)}")
    used = sorted({o for v in direct.values() for o in v})
    sub_a = amp.loc[used, [p_map[c] for c in pcols]]
    sub_i = ic.loc[used, [p_map[c] for c in pcols]]
    log.append(f"Official activities used: {len(used)}; empty amplitudes set to 0: {int(sub_a.isna().sum().sum())}; "
               f"empty IC set to 0: {int(sub_i.isna().sum().sum())} "
               f"(of which on a non-zero amplitude: {int(((sub_a > 0) & sub_i.isna()).sum().sum())})")
    log.append(f"Pressures with demonstrateur = 1 (leaves and nodes): {len(demo)}")
    with open(os.path.join(a.report_dir, "conversion_log.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(log) + "\n")
    print("\n".join(log))
    print(f"Written: {out}")


if __name__ == "__main__":
    main()
